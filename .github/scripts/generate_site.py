#!/usr/bin/env python3
"""Generate a static website from ObsDataPK_OSP.xlsx for GitHub Pages deployment.

Usage: python3 generate_site.py <xlsx_path> <output_dir>
"""

import json
import sys
from pathlib import Path
from datetime import datetime, timezone

import openpyxl

# ---------------------------------------------------------------------------
# Sheet configuration
# ---------------------------------------------------------------------------

SHEETS = ["Studies", "PK-Parameter", "PK-Profiles", "DDI", "Analyte", "Projects"]

SHEET_CONFIG = {
    "Studies": {
        "skip_rows": 1,  # Row 0 has merged group headers; row 1 has column names
        "description": "Overview of all pharmacokinetic studies in the database.",
        "filter_cols": ["Compound/Analyte", "Species", "Route", "Data type"],
        "display_cols": [
            "ID", "Study", "Reference", "Grouping", "Compound/Analyte",
            "Compartment", "Data type", "Species", "Dose", "Dose Unit", "Route", "N",
        ],
        "id_col": "ID",
        "related": [
            {"sheet": "PK-Profiles", "from_col": "ID", "to_col": "ID",
             "label": "PK Profiles"},
            {"sheet": "PK-Parameter", "from_col": "ID", "to_col": "ID",
             "label": "PK Parameters"},
        ],
    },
    "PK-Parameter": {
        "skip_rows": 0,
        "description": "PK parameters (AUC, Cmax, CL) extracted from the studies.",
        "filter_cols": ["Analyte"],
        "display_cols": [
            "ID", "Study", "Reference", "Grouping", "Analyte",
            "AUC Avg", "AUC AvgUnit", "AUC AvgType",
            "Cmax Avg", "Cmax AvgUnit", "Cmax AvgType",
            "CL Avg", "CL AvgUnit",
        ],
        "id_col": "ID",
        "related": [
            {"sheet": "Studies", "from_col": "ID", "to_col": "ID",
             "label": "Study Details"},
        ],
    },
    "PK-Profiles": {
        "skip_rows": 0,
        "description": (
            "Summary of available concentration–time PK profiles "
            "(one row per unique profile; # Time Points shows data density)."
        ),
        "filter_cols": ["Analyte", "Compartment"],
        "display_cols": [
            "ID", "Study", "Reference", "Grouping", "Analyte",
            "Compartment", "# Time Points", "Time range",
        ],
        # Aggregate raw rows into unique profiles
        "aggregate_by": ["ID", "Study", "Reference", "Grouping", "Analyte", "Compartment"],
        "id_col": "ID",
        "related": [
            {"sheet": "Studies", "from_col": "ID", "to_col": "ID",
             "label": "Study Details"},
        ],
    },
    "DDI": {
        "skip_rows": 0,
        "description": "Drug–drug interaction (DDI) data.",
        "filter_cols": ["Victim", "Perpetrator", "Mechanism"],
        "display_cols": [
            "ID", "Study ID", "Reference", "Grouping",
            "Victim", "Perpetrator", "Route Victim", "Route Perpetrator",
            "Mechanism", "Compartment", "AUCR Avg", "CmaxR Avg",
        ],
        "id_col": "ID",
        "related": [],
    },
    "Analyte": {
        "skip_rows": 0,
        "description": "Analyte/compound properties (molecular weight).",
        "filter_cols": ["Name"],
        "display_cols": ["Name", "Molecular weight [g/mol]"],
        "id_col": None,
        "related": [
            {"sheet": "Studies", "from_col": "Name", "to_col": "Compound/Analyte",
             "label": "Studies"},
        ],
    },
    "Projects": {
        "skip_rows": 0,
        "description": "Project assignments for each data series.",
        "filter_cols": ["Analyte", "Projects"],
        "display_cols": [
            "ID", "Study", "Reference", "Grouping", "Analyte",
            "CT profile available", "Projects",
        ],
        "id_col": "ID",
        "related": [
            {"sheet": "Studies", "from_col": "ID", "to_col": "ID",
             "label": "Study Details"},
        ],
    },
}

# ---------------------------------------------------------------------------
# Excel reading
# ---------------------------------------------------------------------------


def _unique_headers(raw_headers):
    """Build a list of unique column names from a raw header row."""
    headers = []
    seen = {}
    for i, h in enumerate(raw_headers):
        name = str(h).strip() if h is not None else f"_col{i}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        headers.append(f"{name}_{count}" if count > 0 else name)
    return headers


def read_sheet(ws, skip_rows=0):
    """Return (headers, list-of-row-dicts) for a worksheet."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = _unique_headers(rows[skip_rows])
    data = []
    for row in rows[skip_rows + 1:]:
        if not any(v is not None for v in row):
            continue
        record = {
            h: (str(row[j]).strip() if row[j] is not None else "")
            for j, h in enumerate(headers)
            if j < len(row)
        }
        data.append(record)
    return headers, data


def aggregate_pk_profiles(rows, group_cols):
    """Collapse PK-Profiles rows into one row per unique profile."""
    seen = {}
    for row in rows:
        key = tuple(row.get(c, "") for c in group_cols)
        if key not in seen:
            seen[key] = {"count": 0, "times": [], **{c: row.get(c, "") for c in group_cols}}
        seen[key]["count"] += 1
        t = row.get("Time", "")
        if t:
            try:
                seen[key]["times"].append(float(t))
            except ValueError:
                pass
    result = []
    for entry in seen.values():
        times = entry.pop("times")
        count = entry.pop("count")
        entry["# Time Points"] = str(count)
        if times:
            entry["Time range"] = f"{min(times):.3g} – {max(times):.3g}"
        else:
            entry["Time range"] = ""
        result.append(entry)
    return result


def read_excel(xlsx_path):
    """Read all configured sheets from the Excel file."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    result = {}
    for sheet_name in SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"Warning: sheet '{sheet_name}' not found, skipping", file=sys.stderr)
            continue
        cfg = SHEET_CONFIG[sheet_name]
        headers, rows = read_sheet(wb[sheet_name], skip_rows=cfg["skip_rows"])
        if "aggregate_by" in cfg:
            rows = aggregate_pk_profiles(rows, cfg["aggregate_by"])
        result[sheet_name] = {"headers": headers, "rows": rows}
        print(f"  {sheet_name}: {len(rows)} rows")
    wb.close()
    return result


# ---------------------------------------------------------------------------
# Data preparation for the site
# ---------------------------------------------------------------------------


def get_unique_values(rows, col):
    """Return sorted unique non-empty values for a column."""
    vals = {
        v.strip()
        for row in rows
        for v in (row.get(col, "") or "").split(",")
        if v.strip() and v.strip() != "None"
    }
    return sorted(vals, key=str.lower)


def prepare_site_data(raw_data):
    """Build sheet_data and filter_values dicts ready for JSON serialisation."""
    sheet_data = {}
    filter_values = {}

    for sheet_name, info in raw_data.items():
        cfg = SHEET_CONFIG[sheet_name]
        # Only keep columns that actually exist
        available = set(info["headers"])
        # For aggregated sheets the new computed cols won't be in headers
        if "aggregate_by" in cfg:
            available.update(["# Time Points", "Time range"])
        display_cols = [c for c in cfg["display_cols"] if c in available]

        rows_out = [{c: row.get(c, "") for c in display_cols} for row in info["rows"]]

        sheet_data[sheet_name] = {
            "cols": display_cols,
            "rows": rows_out,
            "config": {
                "description": cfg["description"],
                "filter_cols": [c for c in cfg["filter_cols"] if c in available],
                "id_col": cfg["id_col"],
                "related": cfg["related"],
            },
        }

        filter_values[sheet_name] = {
            col: get_unique_values(info["rows"], col)
            for col in cfg["filter_cols"]
            if col in available
        }

    return sheet_data, filter_values


# ---------------------------------------------------------------------------
# HTML generation
# ---------------------------------------------------------------------------

HTML_TEMPLATE = """\
<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>OSP PK Database Explorer</title>
  <link rel="stylesheet"
    href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.2/dist/css/bootstrap.min.css">
  <link rel="stylesheet"
    href="https://cdn.datatables.net/2.0.8/css/dataTables.bootstrap5.min.css">
  <style>
    body { background: #f5f6fa; }
    .app-header {
      background: linear-gradient(135deg, #1a3a5c 0%, #2471a3 100%);
      color: #fff; padding: 1.4rem 0; margin-bottom: 1.5rem;
    }
    .app-header h1 { font-size: 1.7rem; margin: 0 0 .2rem; }
    .app-header p  { margin: 0; opacity: .85; font-size: .9rem; }
    .nav-tabs .nav-link         { color: #495057; font-size: .875rem; }
    .nav-tabs .nav-link.active  { font-weight: 600; color: #1a3a5c; }
    .tab-content { background: #fff; border: 1px solid #dee2e6;
                   border-top: 0; border-radius: 0 0 .5rem .5rem; padding: 1rem; }
    .filter-card { border: 1px solid #d0d7de; border-radius: .5rem;
                   background: #f8f9fa; margin-bottom: .75rem; }
    .filter-card .fc-header {
      background: #e2e8f0; border-radius: .5rem .5rem 0 0;
      padding: .45rem .85rem; font-weight: 600; font-size: .85rem;
      display: flex; justify-content: space-between; align-items: center;
    }
    .filter-card .fc-body { padding: .75rem 1rem; }
    .filter-label { font-size: .8rem; font-weight: 500; color: #374151;
                    margin-bottom: .2rem; }
    .active-bar { background: #dbeafe; border: 1px solid #93c5fd;
                  border-radius: .4rem; padding: .4rem .75rem;
                  font-size: .82rem; margin-bottom: .6rem; display: none; }
    .active-bar .tag { display: inline-block; background: #1e40af; color: #fff;
                       border-radius: 3px; padding: .1rem .4rem; margin: .1rem;
                       font-size: .78rem; }
    .btn-rel  { font-size: .75rem; padding: .15rem .45rem; margin: .1rem; }
    .ref-link { max-width: 220px; overflow: hidden; text-overflow: ellipsis;
                white-space: nowrap; display: inline-block; vertical-align: bottom; }
    table.dataTable { font-size: .83rem; }
    .dataTables_wrapper .dataTables_filter input,
    .dataTables_wrapper .dataTables_length select { font-size: .83rem; }
    .sheet-desc { color: #6b7280; font-size: .88rem; margin-bottom: .6rem; }
    .badge-cnt  { font-size: .72rem; }
    footer { font-size: .78rem; color: #9ca3af; text-align: right;
             padding-bottom: .75rem; }
  </style>
</head>
<body>
<header class="app-header">
  <div class="container-fluid px-4">
    <h1>OSP PK Database Explorer</h1>
    <p>Interactive explorer for the Open Systems Pharmacology pharmacokinetic
       observed-data database</p>
  </div>
</header>

<main class="container-fluid px-4">
  <ul class="nav nav-tabs" id="mainTabs" role="tablist"></ul>
  <div class="tab-content" id="mainTabContent"></div>
  <footer>
    Generated %(generated_at)s &nbsp;|&nbsp;
    <a href="https://github.com/pchelle/Database-for-observed-data"
       target="_blank">GitHub repository</a>
  </footer>
</main>

<script src="https://code.jquery.com/jquery-3.7.1.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.2/dist/js/bootstrap.bundle.min.js"></script>
<script src="https://cdn.datatables.net/2.0.8/js/dataTables.min.js"></script>
<script src="https://cdn.datatables.net/2.0.8/js/dataTables.bootstrap5.min.js"></script>
<script>
// ── Embedded data ──────────────────────────────────────────────────────────
const SHEET_DATA    = %(sheet_data_json)s;
const FILTER_VALUES = %(filter_values_json)s;
const SHEETS        = %(sheets_json)s;

// ── Runtime state ──────────────────────────────────────────────────────────
const dts     = {};   // DataTable instances keyed by sheet name
const filters = {};   // activeFilters[sheet] = { col: val, ... }
const inited  = new Set();

// ── Helpers ────────────────────────────────────────────────────────────────
function sid(s)  { return s.replace(/[^\\w]/g, '_'); }
function esc(s)  {
  return String(s ?? '')
    .replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;')
    .replace(/"/g,'&quot;');
}

function renderCell(col, val) {
  if (!val || val === 'None') return '';
  if (col === 'Reference' && /^https?:\\/\\//.test(val)) {
    const short = val.replace(/^https?:\\/\\/(www\\.)?/,'').substring(0,45);
    return `<a href="${esc(val)}" target="_blank" class="ref-link" title="${esc(val)}">`
           + esc(short) + (val.length > 50 ? '…' : '') + '</a>';
  }
  return esc(val);
}

// ── Tab navigation ─────────────────────────────────────────────────────────
function buildNav() {
  const ul = document.getElementById('mainTabs');
  SHEETS.forEach((sheet, i) => {
    if (!SHEET_DATA[sheet]) return;
    const cnt = SHEET_DATA[sheet].rows.length.toLocaleString();
    ul.insertAdjacentHTML('beforeend',
      `<li class="nav-item" role="presentation">
         <a class="nav-link${i===0?' active':''}"
            id="tab-${sid(sheet)}-tab"
            data-bs-toggle="tab"
            href="#tab-${sid(sheet)}"
            role="tab">
           ${esc(sheet)}
           <span class="badge bg-secondary badge-cnt">${cnt}</span>
         </a>
       </li>`);
  });
}

// ── Tab content ────────────────────────────────────────────────────────────
function buildContent() {
  const container = document.getElementById('mainTabContent');
  SHEETS.forEach((sheet, i) => {
    if (!SHEET_DATA[sheet]) return;
    const cfg  = SHEET_DATA[sheet].config;
    const s    = sid(sheet);
    const opts = FILTER_VALUES[sheet] ?? {};

    /* filter dropdowns */
    let filterHtml = '';
    if (cfg.filter_cols.length) {
      const selects = cfg.filter_cols.map(col => {
        const fid = `f-${s}-${sid(col)}`;
        const optsHtml = (opts[col] ?? [])
          .map(v => `<option value="${esc(v)}">${esc(v)}</option>`)
          .join('');
        return `<div class="col-md-3 col-sm-6">
          <div class="filter-label">${esc(col)}</div>
          <select id="${fid}" class="form-select form-select-sm"
            data-sheet="${esc(sheet)}" data-col="${esc(col)}"
            onchange="onFilterChange(this)">
            <option value="">— All —</option>${optsHtml}
          </select></div>`;
      }).join('');
      filterHtml = `
        <div class="filter-card">
          <div class="fc-header">
            <span>&#128269; Filter data</span>
            <button class="btn btn-sm btn-outline-secondary"
              onclick="resetFilters(${JSON.stringify(sheet)})">Reset</button>
          </div>
          <div class="fc-body"><div class="row g-2">${selects}</div></div>
        </div>`;
    }

    /* "view related" column header */
    const hasRel = cfg.related && cfg.related.length > 0;
    const relTh  = hasRel ? '<th>View Related</th>' : '';

    /* column headers */
    const thHtml = SHEET_DATA[sheet].cols
      .map(c => `<th>${esc(c)}</th>`).join('') + relTh;

    container.insertAdjacentHTML('beforeend',
      `<div class="tab-pane fade${i===0?' show active':''}"
            id="tab-${s}" role="tabpanel">
         <p class="sheet-desc">${esc(cfg.description)}</p>
         ${filterHtml}
         <div class="active-bar" id="abar-${s}">
           <strong>Active filters:</strong>
           <span id="atags-${s}"></span>
           <button class="btn btn-sm btn-outline-secondary ms-1"
             onclick="resetFilters(${JSON.stringify(sheet)})">Clear</button>
         </div>
         <div class="table-responsive">
           <table id="dt-${s}"
                  class="table table-striped table-hover table-bordered"
                  style="width:100%">
             <thead><tr>${thHtml}</tr></thead>
             <tbody></tbody>
           </table>
         </div>
       </div>`);

    filters[sheet] = {};
  });
}

// ── DataTable initialisation ───────────────────────────────────────────────
function initTable(sheet) {
  const sdata = SHEET_DATA[sheet];
  const cfg   = sdata.config;
  const s     = sid(sheet);
  const hasRel = cfg.related && cfg.related.length > 0;

  const columns = sdata.cols.map(col => ({
    title : col,
    data  : col,
    render: (data, type) =>
      type === 'display' ? renderCell(col, data) : (data ?? '')
  }));

  if (hasRel) {
    columns.push({
      title    : 'View Related',
      data     : null,
      orderable: false,
      render   : (data, type, row) => {
        if (type !== 'display') return '';
        return cfg.related.map(rel => {
          const val = row[rel.from_col] ?? '';
          if (!val) return '';
          return `<button class="btn btn-outline-primary btn-sm btn-rel"
            onclick="goToSheet(${JSON.stringify(rel.sheet)},
                               ${JSON.stringify(rel.to_col)},
                               ${JSON.stringify(val).replace(/</g,'\\u003c')})"
            >${esc(rel.label)}</button>`;
        }).join('');
      }
    });
  }

  dts[sheet] = $(`#dt-${s}`).DataTable({
    data      : sdata.rows,
    columns   : columns,
    pageLength: 25,
    lengthMenu: [[10,25,50,100],[10,25,50,100]],
    order     : [],
    scrollX   : true,
    autoWidth : false,
    language  : {
      emptyTable : 'No records match the current filters.',
      info       : 'Showing _START_ to _END_ of _TOTAL_ entries',
      infoEmpty  : 'Showing 0 entries',
      infoFiltered: '(filtered from _MAX_ total)',
    }
  });
}

// ── Filtering ──────────────────────────────────────────────────────────────
function onFilterChange(sel) {
  setFilter(sel.dataset.sheet, sel.dataset.col, sel.value);
}

function setFilter(sheet, col, val) {
  filters[sheet] = filters[sheet] ?? {};
  if (val) {
    filters[sheet][col] = val;
  } else {
    delete filters[sheet][col];
  }
  updateFilterBar(sheet);
  applyFilters(sheet);
}

function applyFilters(sheet) {
  const rows    = SHEET_DATA[sheet].rows;
  const active  = filters[sheet] ?? {};
  const entries = Object.entries(active);

  const filtered = entries.length === 0 ? rows : rows.filter(row =>
    entries.every(([col, val]) => {
      const cell = row[col] ?? '';
      return cell === val
          || cell.split(',').map(s => s.trim()).includes(val);
    })
  );

  if (dts[sheet]) {
    dts[sheet].clear().rows.add(filtered).draw();
  }
}

function updateFilterBar(sheet) {
  const s     = sid(sheet);
  const bar   = document.getElementById(`abar-${s}`);
  const tags  = document.getElementById(`atags-${s}`);
  const active = filters[sheet] ?? {};
  const entries = Object.entries(active);

  if (entries.length === 0) {
    bar.style.display = 'none';
    return;
  }
  bar.style.display = 'block';
  tags.innerHTML = entries
    .map(([c,v]) => `<span class="tag">${esc(c)}: ${esc(v)}</span>`)
    .join(' ');
}

function resetFilters(sheet) {
  filters[sheet] = {};
  const cfg = SHEET_DATA[sheet].config;
  const s   = sid(sheet);
  cfg.filter_cols.forEach(col => {
    const el = document.getElementById(`f-${s}-${sid(col)}`);
    if (el) el.value = '';
  });
  updateFilterBar(sheet);
  if (dts[sheet]) {
    dts[sheet].clear().rows.add(SHEET_DATA[sheet].rows).draw();
  }
}

// ── Cross-sheet navigation ─────────────────────────────────────────────────
function goToSheet(targetSheet, col, val) {
  filters[targetSheet] = filters[targetSheet] ?? {};
  filters[targetSheet][col] = val;

  // Update select element if it exists
  const ts  = sid(targetSheet);
  const sel = document.getElementById(`f-${ts}-${sid(col)}`);
  if (sel) sel.value = val;

  // Switch tab
  const tabEl = document.getElementById(`tab-${ts}-tab`);
  if (tabEl) {
    tabEl.click();
    setTimeout(() => {
      updateFilterBar(targetSheet);
      applyFilters(targetSheet);
    }, 80);
  }
}

// ── Bootstrap ──────────────────────────────────────────────────────────────
document.addEventListener('DOMContentLoaded', () => {
  buildNav();
  buildContent();

  // Initialise first visible tab
  const first = SHEETS.find(s => SHEET_DATA[s]);
  if (first) { initTable(first); inited.add(first); }

  // Lazy-init remaining tabs on first show
  document.getElementById('mainTabs').addEventListener('shown.bs.tab', e => {
    const href  = e.target.getAttribute('href');
    const sheet = SHEETS.find(s => `#tab-${sid(s)}` === href);
    if (sheet && !inited.has(sheet)) {
      initTable(sheet);
      inited.add(sheet);
      // Re-apply any pending filters (e.g. from cross-sheet navigation)
      if (Object.keys(filters[sheet] ?? {}).length) {
        updateFilterBar(sheet);
        applyFilters(sheet);
      }
    }
  });
});
</script>
</body>
</html>
"""


def generate_html(sheet_data, filter_values, generated_at):
    html = HTML_TEMPLATE
    html = html.replace("%(generated_at)s", generated_at)
    html = html.replace("%(sheet_data_json)s",
                        json.dumps(sheet_data, ensure_ascii=False))
    html = html.replace("%(filter_values_json)s",
                        json.dumps(filter_values, ensure_ascii=False))
    html = html.replace("%(sheets_json)s", json.dumps(SHEETS))
    return html


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def main():
    if len(sys.argv) != 3:
        print(f"Usage: {sys.argv[0]} <xlsx_path> <output_dir>", file=sys.stderr)
        sys.exit(1)

    xlsx_path  = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])

    if not xlsx_path.exists():
        print(f"Error: '{xlsx_path}' not found", file=sys.stderr)
        sys.exit(1)

    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Reading {xlsx_path} ...")
    raw_data = read_excel(xlsx_path)

    print("Preparing data ...")
    sheet_data, filter_values = prepare_site_data(raw_data)

    print("Generating HTML ...")
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    html = generate_html(sheet_data, filter_values, generated_at)

    out = output_dir / "index.html"
    out.write_text(html, encoding="utf-8")
    print(f"Done — {out}  ({len(html) / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
